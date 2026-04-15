from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"C:\Users\32020\Desktop\occupational _competencies_platform")
SOURCE = ROOT / "系统架构与详细设计说明书-职业能力平台_详细源稿.md"
OUTPUT = ROOT / "系统架构与详细设计说明书-职业能力平台.docx"


def set_run_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_format(paragraph, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_paragraph(document, text="", style=None, align=None, font_name="宋体", font_size=12, bold=False, first_line=True):
    p = document.add_paragraph()
    if style:
        p.style = style
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, font_name, font_size, bold)
    set_paragraph_format(p, first_line=first_line)
    return p


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在打开文档后更新。"
    fld_sep.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def clean_md(text):
    text = text.strip()
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text


def add_md_table(document, rows):
    parsed = []
    for line in rows:
        cells = [clean_md(x.strip()) for x in line.strip().strip("|").split("|")]
        parsed.append(cells)
    if len(parsed) < 2:
        return
    header = parsed[0]
    body = parsed[2:]
    table = document.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    for i, value in enumerate(header):
        p = table.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        set_run_font(run, "宋体", 10.5, True)
    for r, row in enumerate(body, start=1):
        for c, value in enumerate(row):
            p = table.cell(r, c).paragraphs[0]
            run = p.add_run(value)
            set_run_font(run, "宋体", 10.5, False)
    document.add_paragraph()


def prepare_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True


def add_cover(document):
    add_paragraph(document, "卷    号", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "卷内编号", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "密    级", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    for _ in range(4):
        document.add_paragraph()
    add_paragraph(document, "项目编号：OCP20260415SD001", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    document.add_paragraph()
    add_paragraph(document, "职业能力大数据服务平台", align=WD_ALIGN_PARAGRAPH.CENTER, font_name="黑体", font_size=22, bold=True, first_line=False)
    document.add_paragraph()
    add_paragraph(document, "系统架构与详细设计说明书", align=WD_ALIGN_PARAGRAPH.CENTER, font_name="黑体", font_size=20, bold=True, first_line=False)
    add_paragraph(document, "Version: 1.0", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    for _ in range(3):
        document.add_paragraph()
    add_paragraph(document, "项目承担部门：职业能力平台项目组", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "撰 写 人：项目组", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "完成日期：2026-04-15", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "本文档使用部门：项目组  前端开发  后端开发  测试  运维", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "评审负责人：项目负责人", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_paragraph(document, "评审日期：2026-04-15", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    document.add_page_break()


def add_meta(document):
    add_paragraph(document, "文档信息", style="Heading 1", font_name="黑体", font_size=16, bold=True, first_line=False)
    add_paragraph(document, "标题：职业能力大数据服务平台")
    add_paragraph(document, "作者：项目组")
    add_paragraph(document, "创建日期：2026-04-15")
    add_paragraph(document, "上次更新日期：2026-04-15")
    add_paragraph(document, "版本：Version 1.0")
    add_paragraph(document, "部门名称：职业能力平台项目组")
    document.add_paragraph()

    add_paragraph(document, "修订文档历史记录", style="Heading 1", font_name="黑体", font_size=16, bold=True, first_line=False)
    table = document.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    headers = ["日期", "版本", "说明", "作者"]
    values = ["2026-04-15", "V1.0", "基于当前项目实现重写系统架构与详细设计说明书", "项目组"]
    for i, text in enumerate(headers):
        p = table.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, "宋体", 10.5, True)
    for i, text in enumerate(values):
        p = table.cell(1, i).paragraphs[0]
        run = p.add_run(text)
        set_run_font(run, "宋体", 10.5, False)
    document.add_page_break()


def add_contents_page(document):
    add_paragraph(document, "目 录", align=WD_ALIGN_PARAGRAPH.CENTER, font_name="黑体", font_size=16, bold=True, first_line=False)
    p = document.add_paragraph()
    add_toc(p)
    document.add_page_break()


def add_body_from_markdown(document, markdown_text):
    lines = markdown_text.splitlines()
    table_buffer = []
    in_code = False
    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            in_code = not in_code
            continue

        if in_code:
            p = add_paragraph(document, line, font_name="Consolas", font_size=10, first_line=False)
            p.paragraph_format.left_indent = Cm(0.74)
            continue

        if line.strip().startswith("|"):
            table_buffer.append(line)
            continue
        if table_buffer:
            add_md_table(document, table_buffer)
            table_buffer = []

        if not line.strip():
            document.add_paragraph()
            continue

        if line.startswith("# "):
            add_paragraph(document, clean_md(line[2:]), style="Heading 1", font_name="黑体", font_size=16, bold=True, first_line=False)
            continue
        if line.startswith("## "):
            add_paragraph(document, clean_md(line[3:]), style="Heading 2", font_name="黑体", font_size=14, bold=True, first_line=False)
            continue
        if line.startswith("### "):
            add_paragraph(document, clean_md(line[4:]), style="Heading 3", font_name="黑体", font_size=12, bold=True, first_line=False)
            continue
        if re.match(r"^\d+\.\s+", line.strip()):
            add_paragraph(document, clean_md(line.strip()))
            continue
        if line.strip().startswith("- "):
            add_paragraph(document, "• " + clean_md(line.strip()[2:]), first_line=False)
            continue

        add_paragraph(document, clean_md(line))

    if table_buffer:
        add_md_table(document, table_buffer)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(str(SOURCE))

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    prepare_styles(document)
    add_cover(document)
    add_meta(document)
    add_contents_page(document)
    add_body_from_markdown(document, SOURCE.read_text(encoding="utf-8"))
    document.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
