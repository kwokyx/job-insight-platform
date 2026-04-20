from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_PATH = OUTPUT_DIR / "职业能力大数据服务平台-测试计划-完善版.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_document(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)

    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        s = doc.styles[name]
        s.font.name = "黑体"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        s.font.bold = True
    doc.sections[0].top_margin = Cm(2.54)
    doc.sections[0].bottom_margin = Cm(2.54)
    doc.sections[0].left_margin = Cm(3.0)
    doc.sections[0].right_margin = Cm(2.5)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("四川华迪信息技术有限公司")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(16)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("职业能力大数据服务平台")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(22)
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("测试计划（完善版）")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    run.bold = True

    doc.add_paragraph()
    info = [
        ("项目编号", "HDRD20221101"),
        ("撰写人", "郑熙桐"),
        ("完成日期", "2026-04-20"),
        ("评审负责人", "李思源"),
        ("评审日期", "2026-04-20"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for left, right in info:
        row = table.add_row().cells
        set_cell_text(row[0], left, bold=True)
        set_cell_text(row[1], right)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("说明：本版本补充了关键业务接口、可复现前置条件、角色权限测试实例与测试记录字段。")
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(11)

    doc.add_section(WD_SECTION.NEW_PAGE)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(header_cells[idx], header, bold=True)
        set_cell_shading(header_cells[idx], "D9EAF7")
        if widths and idx < len(widths):
            header_cells[idx].width = Cm(widths[idx])
    for row_data in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_data):
            set_cell_text(row[idx], value)
            if widths and idx < len(widths):
                row[idx].width = Cm(widths[idx])
    return table


def add_section_text(doc):
    doc.add_heading("1. 简介", level=1)
    doc.add_heading("1.1 目的", level=2)
    doc.add_paragraph(
        "《职业能力大数据服务平台测试计划》用于明确系统测试阶段的测试范围、测试需求、测试环境、"
        "前置条件、测试账号、关键业务接口、测试记录要求及完成标准，确保测试过程可复现、测试结论可追溯。"
    )
    doc.add_paragraph(
        "本计划特别补充以下老师重点要求：一是验证大数据结果接口是否满足业务需求；二是明确每条测试用例的前置条件、"
        "登录用户/角色、输入数据、预期结果、实际结果、返回码和是否通过。"
    )

    doc.add_heading("1.2 背景", level=2)
    doc.add_paragraph(
        "职业能力大数据服务平台面向学校就业指导老师、学生以及第三方合作系统，围绕行业就业数据开展分布式采集、"
        "数据处理、统计分析、报告生成、岗位推荐、AI辅助问答和开放接口服务，为教学改革、课程优化和学生就业服务提供数据支撑。"
    )

    doc.add_heading("1.3 范围", level=2)
    add_bullets(
        doc,
        [
            "本项目测试定位于系统测试阶段，覆盖计划测试、设计测试、执行测试和评估测试。",
            "覆盖功能测试、接口测试、权限测试、异常测试、边界测试、安全测试、兼容性测试和基础性能测试。",
            "覆盖用户登录/退出、分布式数据采集、数据处理与存储、就业分析、报告自动生成、岗位推送、AI助手分析、对外公开API、平台管理与监控、用户画像与个人中心。",
            "覆盖管理员、教师、学生、第三方调用方四类角色及其权限边界。",
        ],
    )

    doc.add_heading("1.4 定义", level=2)
    add_bullets(
        doc,
        [
            "缺陷级别：1级最严重，4级最轻微。1级为系统崩溃、核心业务不可用、严重安全漏洞；2级为重要功能异常；3级为一般功能问题；4级为界面或提示类轻微问题。",
            "优先级：P1最高，P4最低。存在1级或2级缺陷时，不得将4级问题标记为最高优先级处理。",
            "RBAC：基于角色的访问控制；JWT：基于Token的身份认证；SSE：服务端流式推送；API Key：开放接口鉴权凭证。",
        ],
    )

    doc.add_heading("2. 测试环境与前置条件", level=1)
    doc.add_paragraph("为保证测试过程可复现，执行测试前必须完成以下环境核对。")

    add_table(
        doc,
        ["检查项", "配置/要求", "检查方式"],
        [
            ["前端访问地址", "http://localhost", "浏览器访问首页成功"],
            ["后端服务", "Spring Boot，端口 8080", "访问 http://localhost:8080/doc.html 或接口返回成功"],
            ["算法服务", "FastAPI，端口 8000", "访问 http://localhost:8000/health 成功"],
            ["MySQL", "localhost:3307，库名 career_platform", "数据库连接成功，可查询 sys_user 表"],
            ["Redis", "localhost:6379", "执行 redis-cli ping 返回 PONG"],
            ["RabbitMQ", "5672，管理端 15672", "管理页可访问，消息服务状态正常"],
            ["JWT 配置", "JWT_SECRET 已配置", "登录后可成功签发 accessToken"],
            ["AI 配置", "AI_PROVIDER=siliconflow，模型已配置", "AI问答接口可正常返回或流式输出"],
            ["Docker 环境", "docker compose 可正常运行", "容器状态为 healthy 或 running"],
            ["测试数据", "已导入基础岗位数据、画像数据、报告数据", "关键页面有可查询数据"],
        ],
        widths=[3.0, 7.0, 6.0],
    )

    doc.add_paragraph()
    doc.add_paragraph("如果环境未预置测试账号，应在测试开始前由管理员创建以下账号，并在测试记录中明确写明使用的是哪个账号。")
    add_table(
        doc,
        ["账号", "角色", "建议密码", "用途", "说明"],
        [
            ["admin_test", "管理员", "Admin@123456", "验证后台管理、用户管理、API Key管理、报告发布", "需具备 ROLE_ADMIN"],
            ["teacher_test", "教师", "Teacher@123456", "验证课程管理、教学改革分析、教师权限范围", "需具备 ROLE_TEACHER"],
            ["student_test", "学生", "Student@123456", "验证个人中心、岗位推荐、AI问答、报告生成", "普通用户角色"],
            ["api_test_key", "第三方调用方", "系统生成", "验证对外公开接口、限流、审计日志", "由管理员创建并记录权限范围"],
        ],
        widths=[3.0, 2.0, 3.5, 5.0, 4.5],
    )

    doc.add_heading("3. 关键测试需求", level=1)
    add_table(
        doc,
        ["测试类别", "重点内容", "判定标准"],
        [
            ["业务结果接口", "验证岗位查询、行业快照、分析结果、报告结果、推荐结果是否满足业务需求", "返回字段完整、数据可读、支持教师/学生/第三方使用场景"],
            ["权限控制", "验证管理员、教师、学生是否只能访问各自授权接口", "未授权访问被拒绝，越权访问返回 401/403 或业务拒绝结果"],
            ["异常与边界", "验证空值、超长、非法格式、弱网、服务异常、重复提交", "系统不崩溃，提示清晰，日志可追溯"],
            ["性能与稳定性", "验证登录、常规查询、热点分析、报告提交基础性能", "登录小于 2 秒，查询建议 3 秒内返回，无明显超时或崩溃"],
        ],
        widths=[3.0, 7.0, 7.0],
    )

    doc.add_heading("4. 测试策略", level=1)
    add_bullets(
        doc,
        [
            "功能测试：使用合法和非法输入验证业务流程、字段规则、状态流转和异步任务结果。",
            "接口测试：对登录、报告、推荐、AI问答、Open API 等接口校验参数、结构、状态码、鉴权和返回业务含义。",
            "权限测试：分别以管理员、教师、学生身份访问同一类接口，检查是否出现越权。",
            "异常测试：模拟 MySQL 不可用、Redis 不可用、算法服务超时、断网或弱网场景，检查容错与日志。",
            "边界测试：对文件上传大小、页码 pageSize、空字符串、超长用户名、错误 API Key、错误 Token 等进行验证。",
        ],
    )

    doc.add_heading("5. 测试记录要求", level=1)
    doc.add_paragraph(
        "所有测试执行结果必须记录在测试记录中，且每条记录至少包含：用例编号、前置条件、登录用户/角色、"
        "输入数据、执行步骤、预期结果、实际结果、返回码、是否通过、缺陷编号。"
    )

    doc.add_heading("6. 典型测试计划实例", level=1)
    doc.add_paragraph(
        "以下实例用于说明本项目测试计划应如何写到可执行、可复现的粒度。"
        "其中“实际结果”“返回码”“是否通过”为测试执行后填写项；若已完成联调，可直接按实测填写。"
    )


def add_case(doc, case_id, title, precondition, login_role, input_data, steps, expected, actual, code, passed):
    doc.add_heading(f"{case_id} {title}", level=2)
    add_table(
        doc,
        ["字段", "内容"],
        [
            ["前置条件", precondition],
            ["登录用户/角色", login_role],
            ["输入数据", input_data],
            ["测试步骤", steps],
            ["预期结果", expected],
            ["实际结果", actual],
            ["返回码", code],
            ["是否通过", passed],
        ],
        widths=[3.2, 13.8],
    )
    doc.add_paragraph()


def add_cases(doc):
    add_case(
        doc,
        "TP-AUTH-001",
        "管理员正常登录",
        "前端、后端、MySQL、Redis 均已启动；浏览器可访问 http://localhost；管理员账号 admin_test 已创建且状态正常。",
        "admin_test / 管理员",
        "POST /api/v1/auth/login，username=admin_test，password=Admin@123456",
        "1. 打开登录页。2. 输入管理员账号密码。3. 点击登录。4. 检查是否跳转管理首页并返回 token。",
        "登录成功；返回 accessToken、refreshToken、expiresIn 和用户角色信息；页面展示管理员菜单。",
        "示例填写：登录成功，跳转到后台首页，用户信息中 roleType=1。",
        "HTTP 200，业务码 200",
        "待填写",
    )

    add_case(
        doc,
        "TP-AUTH-002",
        "教师访问管理员接口的越权测试",
        "后端服务已启动；teacher_test 账号已创建；已通过教师账号登录并拿到 JWT。",
        "teacher_test / 教师",
        "GET /api/v1/admin/dashboard，Authorization=Bearer <teacher_token>",
        "1. 使用教师账号登录。2. 复制 JWT。3. 调用管理员仪表盘接口。4. 记录接口响应。",
        "教师用户无权访问管理员接口；系统拒绝访问，不返回管理员数据。",
        "待测试时填写，例如：接口被拒绝，未返回 dashboard 数据。",
        "预期 HTTP 403 或统一错误响应码",
        "待填写",
    )

    add_case(
        doc,
        "TP-AUTH-003",
        "学生访问管理员用户列表的越权测试",
        "后端服务已启动；student_test 账号已创建；学生账号已正常登录。",
        "student_test / 学生",
        "GET /api/v1/admin/users?page=1&pageSize=20",
        "1. 使用学生账号登录。2. 携带 token 访问用户列表接口。3. 观察返回结果。",
        "学生不能查看用户管理数据；系统拒绝访问；不应返回用户列表。",
        "待测试时填写。",
        "预期 HTTP 403 或统一错误响应码",
        "待填写",
    )

    add_case(
        doc,
        "TP-OPEN-001",
        "公开岗位查询接口满足业务需求",
        "后端服务已启动；岗位基础数据已导入；若采用 API Key 模式，api_test_key 已生成并有效。",
        "匿名用户 或 api_test_key / 第三方调用方",
        "GET /api/v1/open/jobs?keyword=Java&city=成都&page=1&pageSize=20",
        "1. 调用公开岗位查询接口。2. 检查分页、岗位标题、企业名称、城市、行业、薪资等字段。3. 评估是否满足外部查询场景。",
        "接口成功返回岗位列表；字段包含 id、title、companyName、city、industryName、education、experience、salaryText、publishDate；pageSize 不超过 50。",
        "待测试时填写，例如：返回 20 条岗位记录，字段齐全，可用于岗位展示与筛选。",
        "HTTP 200，业务码 200",
        "待填写",
    )

    add_case(
        doc,
        "TP-RES-001",
        "行业分析快照接口满足业务需求",
        "后端服务已启动；分析数据已准备；至少存在目标行业和城市样本数据。",
        "匿名用户 或 api_test_key / 第三方调用方",
        "GET /api/v1/open/analysis/industry?industry=人工智能&city=成都",
        "1. 调用行业快照接口。2. 检查是否返回行业分析结构。3. 校验结果能否支持老师查看行业趋势、岗位需求和技能热度。",
        "接口返回行业快照数据；输出内容应可支持业务分析，不应只返回空壳成功信息。",
        "待测试时填写，例如：成功返回行业概览、岗位数量、热门技能、区域分布等信息。",
        "HTTP 200，业务码 200",
        "待填写",
    )

    add_case(
        doc,
        "TP-REPORT-001",
        "学生生成个人分析报告",
        "student_test 已登录；报告服务可用；数据库、算法服务正常；用户画像已填写或存在默认画像。",
        "student_test / 学生",
        "POST /api/v1/reports/generate，reportName=学生综合报告，reportType=COMPREHENSIVE，params={...}",
        "1. 以学生身份登录。2. 提交报告生成请求。3. 记录 taskId。4. 查询 /api/v1/reports/{taskId}/status。5. 验证报告生成结果。",
        "创建报告任务成功；返回 taskId；状态可从 PENDING 变为 SUCCESS；报告在个人报告列表中可见。",
        "待测试时填写，例如：任务创建成功，状态轮询后变为 SUCCESS。",
        "提交时 HTTP 200，查询状态 HTTP 200",
        "待填写",
    )

    add_case(
        doc,
        "TP-TEACHER-001",
        "教师新增课程并进行教学改革分析",
        "teacher_test 已登录；教师账号具备 ROLE_TEACHER；数据库中允许写入课程数据。",
        "teacher_test / 教师",
        "POST /api/v1/teacher/courses，courseName=Java程序设计，coreSkills=Java,SpringBoot,MySQL",
        "1. 教师登录。2. 新增课程。3. 调用 GET /api/v1/teacher/teaching-reform。4. 查看分析结果。",
        "课程添加成功；教学改革分析返回课程与市场需求匹配结果，并生成建议项。",
        "待测试时填写。",
        "HTTP 200，业务码 200",
        "待填写",
    )

    add_case(
        doc,
        "TP-AI-001",
        "学生发起 AI 问答并验证流式返回",
        "student_test 已登录；AI_PROVIDER、AI_API_KEY、AI_MODEL 配置有效；算法服务可用。",
        "student_test / 学生",
        "AI 问题：成都地区 Java 岗位对 Spring Boot 和 MySQL 的需求趋势如何？",
        "1. 登录学生账号。2. 进入 AI 助手页。3. 输入问题并发送。4. 观察 SSE 流式返回。5. 查看历史会话是否保存。",
        "返回内容与问题相关；流式消息持续输出；历史会话可查询；系统无报错中断。",
        "待测试时填写。",
        "HTTP 200 或 SSE 正常建立连接",
        "待填写",
    )

    add_case(
        doc,
        "TP-EXC-001",
        "MySQL 异常下的容错测试",
        "系统正常启动后，手工停止 MySQL 服务或断开数据库连接；记录停止时间与恢复时间。",
        "admin_test / 管理员",
        "调用需要访问数据库的接口，例如 GET /api/v1/admin/users",
        "1. 停止 MySQL。2. 调用后台用户列表接口。3. 观察系统提示和日志。4. 恢复 MySQL 后重新调用。",
        "数据库异常时，系统给出明确错误提示，不应出现前端卡死或无响应；恢复后接口可重新正常使用。",
        "待测试时填写。",
        "异常期返回 5xx 或统一失败码；恢复后返回 HTTP 200",
        "待填写",
    )

    add_case(
        doc,
        "TP-BOUND-001",
        "分页边界值测试",
        "后端服务已启动；公开岗位数据存在。",
        "匿名用户",
        "GET /api/v1/open/jobs?page=1&pageSize=100",
        "1. 调用公开岗位查询接口，pageSize 设为 100。2. 检查服务端是否按最大 50 处理。",
        "系统不应无上限返回；pageSize 应被限制为 50 或按接口规则返回错误提示。",
        "待测试时填写。",
        "HTTP 200，业务码 200，且实际 pageSize 不超过 50；或返回参数错误",
        "待填写",
    )

    doc.add_heading("7. 项目测试活动安排", level=1)
    add_table(
        doc,
        ["阶段", "工作内容", "开始日期", "结束日期", "交付物"],
        [
            ["计划测试", "阅读需求和设计文档，明确接口、角色、环境与风险", "2026-04-18", "2026-04-20", "测试计划"],
            ["设计测试", "设计功能、权限、异常、边界、接口、并发测试用例", "2026-04-18", "2026-04-20", "测试用例"],
            ["执行测试", "按角色和场景执行测试并记录实际结果、返回码和缺陷", "2026-04-21", "2026-04-22", "测试记录、接口测试记录"],
            ["评估测试", "汇总缺陷，分析覆盖率和业务需求达成情况", "2026-04-23", "2026-04-24", "测试日志、测试总结报告"],
        ],
        widths=[2.8, 7.0, 2.5, 2.5, 4.2],
    )

    doc.add_heading("8. 通过准则", level=1)
    add_bullets(
        doc,
        [
            "所有计划测试项均已执行，并完成结果记录。",
            "1级、2级缺陷修复率达到 100%。",
            "3级缺陷修复率达到 80% 以上，4级缺陷完成登记并评估是否延期处理。",
            "关键业务链路“登录 -> 查询/分析 -> 报告/推荐 -> AI问答 -> 管理监控”至少完成一轮完整回归。",
            "关键结果接口经过验证后，能够证明其输出满足教师、学生和第三方调用方的业务需求。",
        ],
    )


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    format_document(doc)
    add_title_page(doc)
    add_section_text(doc)
    add_cases(doc)
    doc.save(str(OUTPUT_PATH))
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
